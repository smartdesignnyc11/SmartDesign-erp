"""reports (BACKLOG §1.13): reporting that is right, proven, and ties to the bank.

Each report is computed by the app from the records and recomputed here, independently, in Python, from the raw
records the page holds (invoices, payments, bills, bank and card lines, time, the GL journals): to the cent. Then the
figures are read off the rendered screens, as the owner, the controller and billing, on the desktop and the phone,
and compared with the recomputation. Parts:

  books     the independent recomputation (Books below) against the app's readers: billed, collected, AR and its
            aging, retainage, the bank and card registers, the GL trial balance and its control accounts, labor,
            each job's cost (bills + bank lines + labor + stock, a bill and the bank line that pays it counted once),
            contract, earned, WIP over/under, backlog, the P&L by project (income, job cost, overhead month by month),
            no-project cost, overhead not on a job, spend not yet coded, the company net; the GL bridge
  screen    the same figures read off P&L Detail, Statements, Monthly Pack, Project Profitability, WIP, Receivables,
            Payables and Cash Forecast: every tie says "Ties", the P&L by project rows equal the recomputation;
            drill-downs open the records with the bank or card line behind each; owner, controller and billing
  cash      the cash basis: income is payments received, cost is what the bank and cards paid; an unpaid payroll run
            is not in it (it was, $16,040); the tie shows what is at the bank and what is not yet deposited
  sealed    the August close: the month's whole P&L and its P&L by project read the same after a late card line and
            the payroll approval; the late line lands in September and P&L Detail lists it with an Open link
  entity    intercompany bills: each entity's column, the eliminations and the consolidated net; P&L by project by
            entity ties
  lines     line income to cost on 121 State: lines + codes with no line + rounding = the job; each line's EAC is
            max(budget, actual + committed); the lines' contract is the job's
  closeout  estimate versus actual by cost code; at close-out a code off by 10% and $500 is proposed to the price
            book (a drafted change naming its Standard); the owner approves it before the cost factor changes
  overhead  overhead by category by month recomputed; break-even monthly revenue = 3 full months' overhead ÷ their
            gross margin
  manual    a manual journal posted by the owner with a reason is flagged on the ties and the Pack, its lines by hand
  cash13    the 13-week forecast: no bill the bank already paid, the owed payroll run and the card balances are out,
            cash today + the weeks' net = the last balance
  excel     the .xlsx the Excel button writes opens (openpyxl) with one sheet per table and numbers as numbers
  ask       Ask answers each report question with the screen's number; the report tables are money, read by a role
            that opens Reports, never by the field or a client
  phone     390px: no sideways scroll on the report tabs, ties one column, tap targets at least 44px

  python3 reports_test.py                 all parts
  python3 reports_test.py books screen    only those parts
All demo data here is invented."""
import asyncio, sys, re, json, base64, io, zipfile
from datetime import datetime
from harness import *

AUG = ['2026-08-01', '2026-08-31']
YTD = ['2026-01-01', '2026-09-03']
EPS = 0.005
COSTED = ('approved', 'paid', 'verified')
ACCRUED = ('submitted', 'pm-approved')

def r2(x):
    # JavaScript's Math.round(100 * x) / 100 (half rounds up), not Python's banker's rounding
    import math
    return math.floor(100 * (x or 0) + 0.5 + 1e-9) / 100

def money(v):
    return ('−$' if v < -EPS else '$') + f'{abs(r2(v)):,.2f}'

DUMP = """() => JSON.stringify({
  DB: {clients: DB.clients, projects: DB.projects.map(p => ({id: p.id, name: p.project.name, clientId: p.clientId, entityId: p.entityId})),
       txns: DB.txns, bills: DB.bills, payroll: DB.payroll, accounts: DB.accounts, pos: DB.pos, time: DB.time,
       employees: DB.employees, entities: DB.entities, gl: DB.gl, stock: DB.stock},
  CAT: CATEGORIES.map(c => ({id: c.id, type: c.type})), today: today()})"""

# ---------------------------------------------------------------------------------------------------------------- the books
class Books:
    """The reports recomputed from the raw records, with nothing of the app's own code."""
    def __init__(self, dump):
        d = json.loads(dump)
        self.D, self.today = d['DB'], d['today']
        self.cat = {c['id']: c['type'] for c in d['CAT']}
        self.clients = self.D['clients']
        self.P = {p['id']: p for p in self.D['projects']}
        self.emp = {e['id']: e for e in self.D['employees']}
        # a bill and the bank or card line that pays it are one cost: the line names the bill (billId), the bill
        # names the line (txnId), or the line carries the matcher's proposal for the bill
        self.pay_of = {}
        for b in self.D['bills']:
            t = next((t for t in self.D['txns'] if t.get('billId') == b['id'] or (b.get('txnId') and t['id'] == b.get('txnId'))
                      or ((t.get('proposal') or {}).get('kind') == 'bill' and (t.get('proposal') or {}).get('billId') == b['id'])), None)
            if t: self.pay_of[b['id']] = t
        self.paid_txn = {t['id']: b for b in self.D['bills'] for t in [self.pay_of.get(b['id'])] if t}

    # documents
    def line_amount(self, l, doc):
        if doc.get('mode') == 'sov':
            st = float(l.get('stored') or 0); sc = float(l.get('scheduled') or 0); pct = float(l.get('pct') or 0)
            st_pct = (100 * st / sc) if st and sc else 0
            if st and pct >= r2(st_pct) - 1e-9: return r2(sc * max(0, pct - r2(st_pct)) / 100 + st)
            return r2(sc * pct / 100)
        ship = r2(float(l.get('shipping') or 0) * (1 + float(l.get('shipMarkupPct') or 0) / 100))
        return r2(r2(float(l.get('qty') or 0) * float(l.get('rate') or 0)) + ship)
    def doc_total(self, d):
        L = d.get('lines') or []
        sub = r2(sum(self.line_amount(l, d) for l in L))
        disc = 0
        if d.get('discount'):
            disc = r2(float(d['discount'].get('value') or 0)) if d['discount'].get('type') == 'amt' else r2(sub * float(d['discount'].get('value') or 0) / 100)
        taxable = r2(sum(self.line_amount(l, d) for l in L if l.get('taxable')))
        tax = r2((taxable - (disc * taxable / sub if sub else 0)) * float(d.get('taxRate') or 0) / 100)
        ret = r2((sub - disc) * float(d.get('retainagePct') or 0) / 100) if d.get('type') == 'invoice' else 0
        dep = r2(float(d.get('depositApplied') or 0)) if d.get('type') == 'invoice' and d.get('mode') == 'sov' else 0
        return {'sub': sub, 'disc': disc, 'tax': tax, 'ret': ret, 'dep': dep, 'total': r2(sub - disc + tax - ret - dep)}
    def doc_net(self, d):
        t = self.doc_total(d); return r2(t['sub'] - t['disc'] - t['dep'])
    def live_invoices(self):
        for c in self.clients:
            for i in c['invoices']:
                if i.get('status') not in ('void', 'draft'): yield c, i
    def invoiced(self, R, pid=None):
        return r2(sum(self.doc_net(i) for c, i in self.live_invoices() if not i.get('retainageRelease') and R[0] <= i['date'] <= R[1] and (pid is None or i.get('projectId') == pid)))
    def collected(self, R):
        return r2(sum(float(p.get('amount') or 0) for c in self.clients for p in c['payments'] if not p.get('unapplied') and R[0] <= p['date'] <= R[1]))
    def ar(self):
        """Open invoices (balance over $0.50, as the client record keeps them) and the cents short-paid below it."""
        open_, short = [], []
        for c, i in self.live_invoices():
            if i.get('retainageRelease'): continue
            paid = sum(float(p.get('amount') or 0) for p in c['payments'] if p.get('invoiceId') == i['id'])
            bal = r2(self.doc_total(i)['total'] - paid)
            age = (datetime.fromisoformat(self.today) - datetime.fromisoformat(i['due'])).days if i.get('due') else 0
            if i.get('status') != 'paid' and bal > 0.5: open_.append((c, i, bal, age))
            elif EPS < bal <= 0.5: short.append((c, i, bal))
        return open_, short
    def aging(self):
        b = {'current': 0, 'd30': 0, 'd60': 0, 'd90': 0}
        for c, i, bal, age in self.ar()[0]:
            k = 'current' if age <= 0 else 'd30' if age <= 30 else 'd60' if age <= 60 else 'd90'
            b[k] = r2(b[k] + bal)
        return b
    def retainage_held(self):
        held = sum(self.doc_total(i)['ret'] for c, i in self.live_invoices())
        released = sum(self.doc_total(i)['total'] for c, i in self.live_invoices() if i.get('retainageRelease'))
        return r2(max(0, held - released))
    # the books
    def gl(self):
        bal = {}
        for j in self.D['gl']['journals']:
            for l in j['lines']:
                bal[l['acct']] = bal.get(l['acct'], 0) + float(l.get('dr') or 0) - float(l.get('cr') or 0)
        return {k: r2(v) for k, v in bal.items()}
    def register(self, a):
        return r2(float(a.get('opening') or 0) + sum(float(t['amt']) for t in self.D['txns'] if t['acct'] == a['id']))
    # cost
    def bill_cost(self, b):
        return 0 if b.get('retainageRelease') else r2(sum(float(l.get('amount') or 0) for l in (b.get('lines') or [])))
    def ctype(self, cat):
        return 'oh' if cat in ('payroll_oh', 'payroll-oh') else self.cat.get(cat or 'uncat', 'non')
    def labor(self, pid, R=None):
        tot = 0
        for t in self.D['time']:
            if t.get('status') == 'rejected' or (R and not (R[0] <= t['date'] <= R[1])): continue
            e = self.emp[t['empId']]; rate = e['rate'] * (1 + e['burden'] / 100)
            w = t.get('work') or []
            if w: tot += sum(x['hours'] * rate for x in w if x['projectId'] == pid)
            elif t.get('projectId') == pid:
                h = (datetime.fromisoformat(t['out']) - datetime.fromisoformat(t['in'])).total_seconds() / 3600 - (t.get('breakMin') or 0) / 60
                tot += h * rate
        return r2(tot)
    def cost_rows(self, view):
        """Every cost, once. view 'pl': as the P&L recognises it (a bill when approved, or when the bank paid it,
        at the earlier date; payroll runs once approved). view 'job': as the job sees it (a bill when entered;
        labor from the time sheets)."""
        rows = []
        for b in self.D['bills']:
            if b.get('status') in ('void', 'rejected') or b.get('ic'): continue
            t = self.pay_of.get(b['id'])
            amt = self.bill_cost(b)
            if t:   # one money: the bill when it is approved (or approved by the PM and paid), else the bank line
                use_bill = b['status'] in COSTED or b['status'] == 'pm-approved'
                rows.append({'kind': 'bill' if use_bill else 'txn', 'date': min(b['date'], t['date']) if use_bill else t['date'], 'amt': amt if use_bill else r2(-float(t['amt'])),
                             'pid': b.get('projectId') or t.get('projectId') or '', 'cat': b.get('cat') or t.get('cat'), 'id': b['id']})
            elif b['status'] in COSTED or (view == 'job' and b['status'] in ACCRUED):
                rows.append({'kind': 'bill', 'date': b['date'], 'amt': amt, 'pid': b.get('projectId') or '', 'cat': b.get('cat'), 'id': b['id']})
        for t in self.D['txns']:
            if t['id'] in self.paid_txn or t.get('billId') or t.get('payrollId'): continue
            if t.get('cat') in ('transfer', 'cc-payment', 'labor', 'income') or (t['amt'] > 0 and self.ctype(t.get('cat')) == 'inc'): continue
            ct = self.ctype(t.get('cat'))
            if ct not in ('cogs', 'oh') and (t.get('cat') or 'uncat') != 'uncat': continue
            rows.append({'kind': 'txn', 'date': t.get('postOn') or t['date'], 'amt': r2(-float(t['amt'])), 'pid': t.get('projectId') or '', 'cat': t.get('cat') or 'uncat', 'id': t['id']})
        if view == 'pl':
            for r in self.D['payroll']:
                if r.get('status') not in COSTED: continue
                for a in r.get('alloc') or []:
                    rows.append({'kind': 'payroll', 'date': r['payDate'], 'amt': r2(float(r['total']) * float(a.get('pct') or 0) / 100), 'pid': a.get('projectId') or '', 'cat': 'labor' if a.get('projectId') else 'payroll_oh', 'id': r['id']})
        return rows
    def job_actual(self, pid, R=None):
        R = R or ['', self.today]
        c = sum(r['amt'] for r in self.cost_rows('job') if r['pid'] == pid and self.ctype(r['cat']) == 'cogs' and R[0] <= r['date'] <= R[1])
        return r2(c + self.labor(pid, R))
    def contract(self, pid):
        c = next(c for c in self.clients if any(p['id'] == pid and p['clientId'] == c['id'] for p in self.D['projects']))
        base = float(((c.get('contracts') or {}).get(pid) or {}).get('net') or 0)
        cos = sum(self.doc_net(x) for x in c['changeOrders'] if x.get('status') in ('approved', 'billed') and x.get('projectId') in (pid, None, ''))
        return r2(base + cos)
    def billed_gross(self, pid):
        return r2(sum(self.doc_net(i) for c, i in self.live_invoices() if i.get('projectId') == pid and not i.get('retainageRelease')))
    # overhead, month by month, shared on the job's cost that month (the last job takes the cents)
    def months(self, R):
        y, m = int(R[0][:4]), int(R[0][5:7]); out = []
        while f'{y}-{m:02d}' <= R[1][:7]:
            out.append(f'{y}-{m:02d}'); m += 1
            if m == 13: y, m = y + 1, 1
        return out
    def month_range(self, ym):
        import calendar
        end = f'{ym}-{calendar.monthrange(int(ym[:4]), int(ym[5:7]))[1]:02d}'
        return [ym + '-01', min(end, self.today)]
    def pl_month(self, ym, pids):
        R = self.month_range(ym)
        rows = [r for r in self.cost_rows('pl') if R[0] <= r['date'] <= R[1]]
        pool = r2(sum(r['amt'] for r in rows if self.ctype(r['cat']) == 'oh'))
        uncoded = r2(sum(r['amt'] for r in rows if r['cat'] == 'uncat'))
        base = {p: self.job_actual(p, R) for p in pids}
        base = {p: v for p, v in base.items() if v > EPS}
        T = sum(base.values()); alloc = {}; left = pool
        for k, p in enumerate(base):
            v = r2(left) if k == len(base) - 1 else r2(pool * base[p] / T)
            alloc[p] = v; left = r2(left - v)
        return {'pool': pool, 'uncoded': uncoded, 'alloc': alloc, 'unalloc': 0 if T else pool, 'cogs': r2(sum(r['amt'] for r in rows if self.ctype(r['cat']) == 'cogs')), 'rows': rows}
    def pl_by_project(self, R, eac):
        """The company P&L for R split by project. eac: each project's estimate at completion (the app's, checked
        separately against its lines); earned = contract × job cost ÷ EAC."""
        pids = [p['id'] for p in self.D['projects']]
        rows = {p: {'income': 0.0, 'cost': 0.0, 'oh': 0.0} for p in pids}
        for p in pids:
            con = self.contract(p) if any(p in (c.get('contracts') or {}) for c in self.clients) else 0
            if con and eac.get(p):
                earned = r2(con * self.job_actual(p) / eac[p])
                # YTD from the year's start: nothing billed or earned before it, so income = invoiced + (earned - billed)
                rows[p]['income'] = r2(self.invoiced(R, p) + earned - self.billed_gross(p))
            else:
                rows[p]['income'] = self.invoiced(R, p)
        noproj = 0; unalloc = 0; uncoded = 0
        for ym in self.months(R):
            M = self.pl_month(ym, pids)
            for p, v in M['alloc'].items(): rows[p]['oh'] = r2(rows[p]['oh'] + v)
            unalloc = r2(unalloc + M['unalloc']); uncoded = r2(uncoded + M['uncoded'])
            for r in M['rows']:
                if self.ctype(r['cat']) != 'cogs': continue
                if r['pid']: rows[r['pid']]['cost'] = r2(rows[r['pid']]['cost'] + r['amt'])
                else: noproj = r2(noproj + r['amt'])
        for v in rows.values(): v['net'] = r2(v['income'] - v['cost'] - v['oh'])
        net = r2(sum(v['net'] for v in rows.values()) - noproj - unalloc - uncoded)
        return {'rows': rows, 'noproj': noproj, 'unalloc': unalloc, 'uncoded': uncoded, 'net': net}
    def cash_pl(self, R):
        income = self.collected(R)
        cogs = oh = 0
        for t in self.D['txns']:
            if not (R[0] <= t['date'] <= R[1]) or t.get('payrollId') or t.get('cat') in ('transfer', 'cc-payment', 'labor', 'income'): continue
            ct = self.ctype(t.get('cat'))
            if ct == 'cogs': cogs += -float(t['amt'])
            elif ct == 'oh' or (t.get('cat') or 'uncat') == 'uncat': oh += -float(t['amt'])
        for r in self.D['payroll']:
            if r.get('status') in ('paid', 'verified') and R[0] <= r['payDate'] <= R[1]:
                for a in r.get('alloc') or []:
                    (cogs, oh) = (cogs + r2(r['total'] * a['pct'] / 100), oh) if a.get('projectId') else (cogs, oh + r2(r['total'] * a['pct'] / 100))
        for b in self.D['bills']:
            if b['id'] not in self.pay_of and b.get('paidAt') and R[0] <= b['paidAt'] <= R[1]:
                if self.ctype(b.get('cat')) == 'cogs': cogs += self.bill_cost(b)
                else: oh += self.bill_cost(b)
        return {'income': r2(income), 'cogsT': r2(cogs), 'ohT': r2(oh), 'net': r2(income - cogs - oh)}

# ---------------------------------------------------------------------------------------------------------------- page helpers
async def page(p, uid='u_owner', phone=False):
    b, pg, errs = await open_app(p, phone=phone)
    await sign_in(pg, uid)
    return b, pg, errs

async def books(pg):
    return Books(await pg.evaluate(DUMP))

async def tab(pg, t, preset=None, proj=None):
    await pg.evaluate("([t, preset, proj]) => { if (preset) REP.preset = preset; if (proj) repProj = proj; repTab = t; view = 'reports'; render(); }", [t, preset, proj])
    await pg.wait_for_timeout(150)
    return await pg.evaluate("() => document.getElementById('rep-print-area') ? document.getElementById('rep-print-area').innerText : ''")

async def screen_ties(pg, sel='.rp-ties'):
    return await pg.evaluate("""s => [...document.querySelectorAll(s + ' .rp-tie')].map(e => ({id: e.dataset.rptie,
        text: e.innerText, ok: !!e.querySelector('.rp-tie-hd .rp-ok'), off: !!e.querySelector('.rp-tie-hd .rp-off'), manual: e.classList.contains('manual')}))""", sel)

def num(s):
    s = s.strip().replace('−', '-').replace('$', '').replace(',', '')
    neg = s.startswith('(') or s.startswith('-')
    s = s.strip('()-')
    try: v = float(s)
    except ValueError: return None
    return -v if neg else v

LIB = re.search(r'LIB = r"""(.*?)"""', open(os.path.join(os.path.dirname(os.path.abspath(__file__)), 'closetie_test.py'), encoding='utf-8').read(), re.S).group(1)

# ---------------------------------------------------------------------------------------------------------------- parts
async def part_books(p, c):
    b, pg, errs = await page(p)
    B = await books(pg)
    A = await pg.evaluate("""() => { const R = ['2026-01-01', today()], P = companyPL(R), Bal = companyBalance(),
        S = rpPL(R), G = glBalances({entityId: 'all', to: today()}), J = {};
        DB.projects.forEach(p => { const x = jobFin(p); J[p.id] = {contract: x.contract, actual: x.actual, eac: x.eac, earned: x.earned, billed: x.billed, overUnder: x.overUnder, backlog: x.backlog, labor: x.laborCost,
            lines: (x.eacLines || []).map(l => ({key: l.key, budget: l.budget, actual: l.actual, committed: l.committed, eac: l.eac}))} });
        const bk = DB.clients.map(clientFin).reduce((a, f) => (['current','d30','d60','d90'].forEach(k => a[k] = (a[k] || 0) + (+f.buckets[k] || 0)), a), {});
        return {P: {income: P.income, invoiced: P.invoiced, wipAdj: P.wipAdj, collected: P.collected, cogsT: P.cogsT, ohT: P.ohT, net: P.net, needs: (P.oh || {})['Needs Category'] || 0},
          B: {ar: Bal.arOut, ret: Bal.ret, wip: Bal.wip, backlog: Bal.backlog, cash: Bal.cash, ap: Bal.ap}, bk,
          S: {net: S.net, ok: S.ok, unalloc: S.unalloc, uncoded: S.uncoded, noproj: S.none.cost, rows: S.rows.map(r => ({id: r.id, income: r.income, cost: r.cost, oh: r.oh, net: r.net}))},
          G: Object.fromEntries(Object.entries(G).map(([k, v]) => [k, v.bal])), J,
          accts: DB.accounts.map(a => ({id: a.id, balance: a.balance, type: a.type})),
          ties: rpTies(R).map(t => ({id: t.id, ok: t.ok, diff: t.diff, manual: !!t.manual})), today: today()} }""")
    P, Bal, S, G, J = A['P'], A['B'], A['S'], A['G'], A['J']
    eq = lambda a, b, msg: c.ok(abs(r2(a) - r2(b)) <= EPS, f'{msg}: app {money(a)} vs recomputed {money(b)}')
    # billed, collected
    eq(P['invoiced'], B.invoiced(YTD), 'invoiced YTD')
    eq(P['collected'], B.collected(YTD), 'collected YTD')
    # AR, aging, the cents short-paid, retainage
    op, short = B.ar()
    eq(Bal['ar'], sum(x[2] for x in op), 'receivables = open invoices')
    for k, v in B.aging().items(): eq(A['bk'].get(k, 0), v, f'AR aging bucket {k}')
    eq(G.get('1100', 0), sum(x[2] for x in op) + sum(x[2] for x in short), 'ledger 1100 = open invoices + short-paid cents')
    c.ok([x[1]['number'] for x in short] == ['INV-2026-035'] and abs(short[0][2] - 0.16) < EPS, f'short-paid invoices {[(x[1]["number"], x[2]) for x in short]} (want INV-2026-035 $0.16)')
    eq(Bal['ret'], B.retainage_held(), 'retainage receivable')
    eq(G.get('1150', 0), B.retainage_held(), 'ledger 1150 retainage receivable')
    # the books balance, and 2400 holds only the deposits still held (the applied retainer is released)
    gl = B.gl()
    c.ok(abs(sum(gl.values())) <= EPS, f'trial balance off by {sum(gl.values())}')
    eq(-gl.get('2400', 0), 24000, 'ledger 2400 customer deposits (the $5,000 45 Cranberry retainer applied on INV-2026-D45-1 is released)')
    # registers
    for a in B.D['accounts']:
        eq(a['balance'], B.register(a), f'{a["name"]} register = opening + lines')
        eq(a['balance'], gl.get('acct:' + a['id'], 0), f'{a["name"]} = its ledger account')
    # labor, job cost, contract, earned, WIP, backlog
    for pid, x in J.items():
        eq(x['labor'], B.labor(pid), f'{pid} labor (clocked hours × rate × (1 + burden))')
        eq(x['actual'], B.job_actual(pid), f'{pid} job cost (bills + bank lines, each money once, + labor + stock)')
        if not x['contract']: continue
        eq(x['contract'], B.contract(pid), f'{pid} contract = base + approved COs')
        if x['lines']:
            for l in x['lines']:
                c.ok(abs(l['eac'] - max(l['budget'], r2(l['actual'] + l['committed']))) <= EPS, f'{pid} line {l["key"]} EAC {l["eac"]} is not max(budget {l["budget"]}, actual + committed {r2(l["actual"] + l["committed"])})')
            eq(x['eac'], max(r2(sum(l['eac'] for l in x['lines'])), r2(x['actual'])), f'{pid} EAC = the lines\' EAC')
        eq(x['earned'], r2(B.contract(pid) * B.job_actual(pid) / x['eac']), f'{pid} earned = contract × cost ÷ EAC')
        eq(x['overUnder'], r2(B.billed_gross(pid) - x['earned']), f'{pid} WIP over/under = billed − earned')
        eq(x['backlog'], r2(B.contract(pid) - x['earned']), f'{pid} backlog = contract − earned')
    eq(Bal['wip'], sum(B.billed_gross(pid) - x['earned'] for pid, x in J.items() if x['contract']), 'net over-billing (WIP)')
    # P&L by project, independently, and the company net it adds up to
    PB = B.pl_by_project(YTD, {pid: x['eac'] for pid, x in J.items()})
    for r in S['rows']:
        mine = PB['rows'][r['id']]
        for k in ('income', 'cost', 'oh', 'net'): eq(r[k], mine[k], f'P&L by project {r["id"]} {k}')
    eq(S['noproj'], PB['noproj'], 'job cost on no project')
    eq(S['unalloc'], PB['unalloc'], 'overhead not on a job (months with no job cost)')
    eq(S['uncoded'], PB['uncoded'], 'spend not yet coded')
    eq(P['net'], PB['net'], 'company net YTD = P&L by project + no project + overhead not on a job + not yet coded')
    eq(S['net'], P['net'], 'rpPL net = companyPL net')
    eq(P['cogsT'], r2(sum(v['cost'] for v in PB['rows'].values()) + PB['noproj']), 'P&L job cost = projects + no project')
    # every month's cost recomputed
    for ym in ('2026-06', '2026-07', '2026-08', '2026-09'):
        M = B.pl_month(ym, list(J))
        Pm = await pg.evaluate("R => { const P = companyPL(R); return {cogsT: P.cogsT, ohT: P.ohT, needs: (P.oh || {})['Needs Category'] || 0} }", B.month_range(ym))
        eq(Pm['cogsT'], M['cogs'], f'{ym} job cost')
        eq(Pm['ohT'], r2(M['pool'] + M['uncoded']), f'{ym} overhead (coded + not yet coded)')
    # the ledger bridge: management cost = ledger cost + vendor prepayments + spend in clearing
    led_cost = r2(sum(v for k, v in gl.items() if k[:1] in '5678'))
    eq(P['cogsT'] + P['ohT'], led_cost + gl.get('1300', 0) + gl.get('1090', 0), 'management cost = ledger cost + prepayments (1300) + clearing (1090)')
    led_rev = r2(-sum(v for k, v in gl.items() if k[:1] == '4'))
    eq(P['invoiced'], led_rev + (-gl.get('2400', 0)), 'invoiced = ledger revenue + deposits still held (2400)')
    # every tie the app computes ties
    off = [t for t in A['ties'] if not t['ok'] and not t['manual']]
    c.ok(not off, f'ties that do not tie: {off}')
    c.ok(len(A['ties']) >= 15, f'only {len(A["ties"])} ties')
    c.ok(not errs, f'books errors {errs[:3]}')
    await b.close()

async def part_screen(p, c):
    for uid in ('u_owner', 'u_controller', 'u_billing'):
        b, pg, errs = await page(p, uid)
        can = await pg.evaluate("(ROLE_VIEWS[role()] || []).includes('reports')")
        c.ok(can, f'{uid} cannot open Reports')
        if not can: await b.close(); continue
        B = await books(pg)
        S = await pg.evaluate("() => { const S = rpPL(rpRange()); return {net: S.net, rows: S.rows.map(r => ({id: r.id, name: rpProjName(r.id), income: r.income, cost: r.cost, oh: r.oh, net: r.net}))} }")
        # P&L Detail: the P&L by Project table, read off the screen
        t = await tab(pg, 'company', 'ty')
        rows = await pg.evaluate("""() => { const T = document.querySelector('[data-rppl] table'); return T ? [...T.rows].map(r => [...r.cells].map(x => x.innerText.trim())) : [] }""")
        c.ok(len(rows) > 2, f'{uid}: no P&L by Project table on P&L Detail')
        for r in S['rows']:
            line = next((x for x in rows if x and x[0] == r['name']), None)
            c.ok(line is not None, f'{uid}: {r["name"]} not on P&L by Project')
            if line:
                got = [num(v) for v in line[2:7]]
                want = [r['income'], r['cost'], r2(r['income'] - r['cost']), r['oh'], r['net']]
                c.ok(all(g is not None and abs(g - w) <= EPS for g, w in zip(got, want)), f'{uid}: {r["name"]} on screen {line[2:7]} vs {want}')
        tot = next((x for x in rows if x and x[0] == 'Company net profit'), None)
        c.ok(tot and abs(num(tot[-2]) - S['net']) <= EPS, f'{uid}: company net on screen {tot} vs {S["net"]}')
        c.ok(re.search(r'Ties\s+Projects\' net', t) is not None, f'{uid}: the P&L by project tie line does not say Ties')
        c.ok('P&L by Client' in t or 'P&L BY CLIENT' in t.upper(), f'{uid}: no P&L by Client')
        # every tie on every tab says Ties
        for tb, sel, at_least in (('company', '[data-rpties=pl]', 2), ('financials', '[data-rpties=all]', 15), ('pack', '[data-rpties=pack]', 15), ('wip', '[data-rpties=wip]', 1),
                                  ('ar', '[data-rpties=ar]', 3), ('payables', '[data-rpties=ap]', 1), ('cash13', '[data-rpties=cash]', 2), ('profit', '[data-rpties=job]', 1)):
            await tab(pg, tb, 'ty', 'p121')
            T = await screen_ties(pg, sel)
            bad = [x['id'] for x in T if not x['manual'] and not x['ok']]
            c.ok(len(T) >= at_least and not bad, f'{uid}: {tb} ties {len(T)} (want ≥{at_least}), not tying: {bad}')
        # drill: the P&L's no-project cost opens the record with its bank line
        await tab(pg, 'company', 'ty')
        await pg.evaluate("() => document.querySelector('[data-rpdrill=\"pl:noproj-cost\"]').click()")
        await pg.wait_for_timeout(200)
        m = await pg.evaluate("() => { const m = document.getElementById('mbm'); return m ? m.innerText : '' }")
        c.ok('ZELLE TO M RODRIGUEZ' in m and 'bank line' in m.lower() and 'statement' in m and '$6,800.00' in m and 'Close' in m, f'{uid}: the no-project drill does not show the Zelle line and its bank backing: {m[:300]!r}')
        # a typed line with no statement behind it is marked as such
        await pg.evaluate("() => { mbCloseModal(); document.querySelector('[data-rpdrill=\"pl:oh-uncoded\"]').click() }")
        await pg.wait_for_timeout(200)
        m = await pg.evaluate("() => { const m = document.getElementById('mbm'); return m ? m.innerText : '' }")
        c.ok(re.search(r'9 with no bank or card line behind them', m) and re.search(r'TINO FLOORING[^\n]*\n+\s*By Hand\s*\n\s*typed in, not yet seen', m, re.I), f'{uid}: an uncoded typed line is not marked as having no bank line: {m[:300]!r}')
        await pg.evaluate("() => mbCloseModal()")
        # Project Profitability: By Line is income to cost and ties to the job
        t = await tab(pg, 'profit', 'ty', 'p121')
        c.ok('BY LINE: INCOME TO COST' in t.upper() and re.search(r'Ties\s+Lines \$95,617\.91', t), f'{uid}: By Line is not income to cost tied to the job: {t[t.upper().find("BY LINE"):][:200]!r}')
        c.ok('$102,670' not in t, f'{uid}: the old double-counted by-line total is still on screen')
        c.ok('Unlinked job cost' not in t, f'{uid}: the by-code table still repeats the bank-line total as unlinked cost')
        c.ok(not errs, f'{uid} screen errors {errs[:3]}')
        await b.close()

async def part_cash(p, c):
    b, pg, errs = await page(p)
    B = await books(pg)
    await pg.evaluate("() => document.querySelector('[data-basis=cash]') || true")
    await tab(pg, 'company', 'ty')
    await pg.evaluate("() => { const x = document.querySelector('[data-basis=cash]'); x ? x.click() : (DB.company.basis = 'cash', render()) }")
    await pg.wait_for_timeout(200)
    A = await pg.evaluate("() => { const P = companyPL(['2026-01-01', today()]); const S = rpPL(['2026-01-01', today()]); const T = rpTies(['2026-01-01', today()], ['pl']); return {basis: basis(), income: P.income, cogsT: P.cogsT, ohT: P.ohT, net: P.net, oh: P.oh, cogs: P.cogs, S: {ok: S.ok, net: S.net}, T: T.map(t => ({id: t.id, ok: t.ok, parts: t.parts}))} }")
    C = B.cash_pl(YTD)
    c.ok(A['basis'] == 'cash', f'basis {A["basis"]}')
    for k in ('income', 'cogsT', 'ohT', 'net'):
        c.ok(abs(A[k] - C[k]) <= EPS, f'cash basis {k}: app {money(A[k])} vs recomputed {money(C[k])}')
    c.ok('Unallocated Labor' not in (A['oh'] or {}) and not any('Labor' in k for k in (A['cogs'] or {})), f'cash basis still counts the unpaid payroll run: {A["oh"]} {A["cogs"]}')
    c.ok(A['S']['ok'], 'cash basis: P&L by project does not tie')
    inc = next(t for t in A['T'] if t['id'] == 'income')
    c.ok(inc['ok'] and abs(inc['parts'][0]['amt'] - 169750) < EPS and abs(inc['parts'][1]['amt'] - 21350) < EPS, f'cash income tie: {inc}')
    t = await pg.evaluate("() => document.getElementById('rep-print-area').innerText")
    c.ok('Cash basis' in t, 'the P&L by project does not say it is cash basis')
    # the payroll run, once paid, is in the cash basis exactly once
    r = await pg.evaluate("""() => { const pr = DB.payroll.find(x => x.id === 'pr1'); pr.status = 'paid'; pr.paidAt = today();
        DB.txns.push({id: 'tpaytest', date: pr.payDate, acct: 'ac1', desc: 'PAYROLL 8/17 – 8/28', amt: -pr.total, cat: 'payroll-oh', projectId: '', vendor: 'Payroll', status: 'categorized', payrollId: pr.id});
        const P = companyPL(['2026-01-01', today()]); return {net: P.net} }""")
    c.ok(abs(r['net'] - r2(C['net'] - 16040)) <= EPS, f'a paid payroll run counts once on the cash basis: net {r["net"]} want {r2(C["net"] - 16040)}')
    await pg.evaluate("() => { DB.company.basis = 'accrual'; save(); }")
    c.ok(not errs, f'cash errors {errs[:3]}')
    await b.close()

async def part_sealed(p, c):
    b, pg, errs = await page(p)
    await pg.evaluate(LIB)
    r = await pg.evaluate("""async () => { await CT.monthEnd(); const x = fin2RunClose('2026-08', 'September lines are coded next month'); return {ok: x.ok, why: x.why || ''} }""")
    c.ok(r['ok'], f'August close: {r}')
    snap_js = """() => { const A = ['2026-08-01', '2026-08-31'], P = companyPL(A), S = rpPL(A), G = glPL('all', A[0], A[1]);
        return {P: {income: P.income, cogsT: P.cogsT, ohT: P.ohT, net: P.net, closed: !!P.closedRead}, G: {rev: G.revT, net: G.net}, S: {net: S.net, ok: S.ok, rows: S.rows.map(r => [r.id, r.income, r.cost, r.oh, r.net])}} }"""
    s0 = await pg.evaluate(snap_js)
    c.ok(s0['P']['closed'], 'August does not read from its seal')
    c.ok(abs(s0['P']['income'] - s0['G']['rev']) <= EPS and abs(s0['P']['net'] - s0['G']['net']) <= EPS, f"closed August: management and ledger disagree {s0['P']} {s0['G']}")
    c.ok(s0['S']['ok'], f'closed August: P&L by project does not tie {s0["S"]}')
    t0 = await tab(pg, 'company', 'lm')
    # late changes while August is closed: a card line dated in August, the payroll run approved
    r = await pg.evaluate("""async () => {
        const late = [{date: '2026-08-28', desc: 'NEMO TILE SUPPLY LATE POST', amt: -1000}];
        await CT.importCsv('ac2', 'Amex_Activity_Sep2026.csv', late, '2026-08-28', '2026-09-03', null);
        const t = DB.txns.find(x => x.desc === 'NEMO TILE SUPPLY LATE POST'); t.cat = 'materials'; t.projectId = 'p121'; t.code = '08'; t.status = 'categorized'; save();
        const pr = DB.payroll.find(x => x.id === 'pr1'); payrollAction(pr, 'approve', 'office staff are salaried');
        return {postOn: t.postOn, pr: pr.status} }""")
    c.ok(r['postOn'] == '2026-09-01', f'late line posting period {r}')
    s1 = await pg.evaluate(snap_js)
    c.ok(json.dumps(s1['P']) == json.dumps(s0['P']), f'sealed August P&L moved: {s0["P"]} -> {s1["P"]}')
    c.ok(json.dumps(s1['S']) == json.dumps(s0['S']), f'sealed August P&L by project moved: {s0["S"]} -> {s1["S"]}')
    t1 = await tab(pg, 'company', 'lm')
    strip = lambda s: re.sub(r'prepared [^\n]*', '', s)
    c.ok(strip(t1.split('Corrections to Sealed Months')[0]) == strip(t0.split('Corrections to Sealed Months')[0]) or 'P&L BY PROJECT' in t1.upper(), 'August P&L Detail reads differently after the late changes')
    # September carries the late line and lists it as a correction, with Open
    t2 = await tab(pg, 'company', 'tm')
    c.ok('CORRECTIONS TO SEALED MONTHS' in t2.upper() and 'NEMO TILE SUPPLY LATE POST' in t2, f'September does not list the correction to sealed August: {t2[:200]!r}')
    s9 = await pg.evaluate("() => { const S = rpPL(['2026-09-01', today()]); return {ok: S.ok, late: rpLate(['2026-09-01', today()]).map(r => rpRowAmt(r))} }")
    c.ok(s9['ok'] and s9['late'] == [1000], f'September with the late line: {s9}')
    await pg.evaluate("() => document.querySelector('[data-rplateopen=\"0\"]').click()")
    await pg.wait_for_timeout(250)
    v = await pg.evaluate('view')
    c.ok(v == 'bank', f'the correction\'s Open went to {v}, not Banking')
    c.ok(not errs, f'sealed errors {errs[:3]}')
    await b.close()

async def part_entity(p, c):
    b, pg, errs = await page(p)
    r = await pg.evaluate("""() => { const R = ['2026-01-01', today()], before = companyPL(R).net, un = icUnbilled();
        icCreateBills(); const E = entityPL(R), after = companyPL(R).net, cons = E.C;
        const inc = E.cols.reduce((a, x) => a + x.P.income, 0), cost = E.cols.reduce((a, x) => a + x.P.cogsT, 0);
        const by = DB.entities.map(e => { repEntity = e.id; const S = rpPL(R), T = rpTies(R, ['pl']); repEntity = ''; return {e: e.id, ok: S.ok, ties: T.every(t => t.ok), net: S.net} });
        return {n: un.length, amt: un.reduce((a, o) => a + o.amount, 0), ic: DB.bills.filter(b => b.ic).length, before, after, elim: E.elim, elimWip: E.elimWip, colsNet: r2(E.cols.reduce((a, x) => a + x.P.net, 0)), consNet: cons.net, inc: r2(inc), cost: r2(cost), consInc: cons.income, consCost: cons.cogsT, icRev: E.cols.map(x => x.P.icRevenue), icCost: E.cols.map(x => x.P.icCost), by} }""")
    c.ok(r['ic'] >= 1, f'no intercompany bill created: {r}')
    # the group billing itself makes no consolidated profit: the markup is eliminated (a cent per bill of rounding)
    c.ok(abs(r['after'] - r['before']) <= 0.01 * r['ic'], f"consolidated net moved {r['before']} -> {r['after']} when the subsidiaries billed their hours")
    c.ok(abs(r2(r['inc'] - r['elim'] - r['elimWip']) - r['consInc']) <= EPS, f"entities' income {r['inc']} − eliminations {r['elim']} − intercompany profit in WIP {r['elimWip']} ≠ consolidated {r['consInc']}")
    c.ok(abs(r2(r['colsNet'] - r['elimWip']) - r['consNet']) <= EPS, f"entities' net {r['colsNet']} − intercompany profit in WIP {r['elimWip']} ≠ consolidated {r['consNet']}")
    c.ok(r['elimWip'] > 0, f"no intercompany profit in WIP eliminated: {r['elimWip']}")
    c.ok(abs(sum(r['icRev']) - r['elim']) <= EPS, f"eliminations {r['elim']} ≠ intercompany revenue {r['icRev']}")
    c.ok(all(x['ok'] and x['ties'] for x in r['by']), f'P&L by project by entity: {r["by"]}')
    c.ok(not errs, f'entity errors {errs[:3]}')
    await b.close()

async def part_lines(p, c):
    b, pg, errs = await page(p)
    r = await pg.evaluate("""() => { const p = DB.projects.find(x => x.id === 'p121'), J = jobFin(p), L = rpLineRows(p), C = rpJobCost(p);
        return {contract: J.contract, actual: J.actual, eac: J.eac, earned: J.earned, rows: L, lineSum: C.lineSum, noLine: C.noLine, rounding: C.lineRounding, lineDiff: C.lineDiff,
          codes: C.codes, legacy: projectProfit(p).lines.reduce((a, l) => a + l.actual, 0)} }""")
    sm = lambda k: r2(sum(x[k] for x in r['rows']))
    c.ok(abs(sm('contract') - r['contract']) <= EPS, f"lines' contract {sm('contract')} vs the job's {r['contract']}")
    c.ok(abs(r2(r['lineSum'] + r['noLine'] + r['rounding']) - r['actual']) <= EPS and r['lineDiff'] == 0, f"lines {r['lineSum']} + no line {r['noLine']} + rounding {r['rounding']} vs the job {r['actual']}")
    c.ok(abs(r['rounding']) <= 0.01 * len(r['rows']), f"rounding {r['rounding']} is more than a cent a line")
    c.ok(abs(sm('eac') - r['eac']) <= EPS, f"lines' EAC {sm('eac')} vs the job's {r['eac']}")
    for x in r['rows']:
        c.ok(abs(x['eac'] - max(x['budget'], r2(x['actual'] + x['committed']))) <= EPS, f"{x['name']}: EAC {x['eac']} ≠ max(budget, actual + committed)")
        c.ok(abs(x['marginAtCompletion'] - r2(x['contract'] - x['eac'])) <= EPS and abs(x['marginToDate'] - r2(x['earned'] - x['actual'])) <= EPS, f"{x['name']}: margins")
        want = r2(x['contract'] * min(1, x['actual'] / x['eac'])) if x['eac'] else 0
        c.ok(abs(x['earned'] - want) <= EPS, f"{x['name']}: earned {x['earned']} vs contract × cost ÷ EAC {want}")
    co = [x for x in r['rows'] if x['co']]
    c.ok(len(co) == 3 and abs(sum(x['contract'] for x in co) - 6200) <= EPS, f'approved CO lines on the report: {[(x["name"], x["contract"]) for x in co]}')
    plumb = next(x for x in r['rows'] if x['name'].startswith('Plumbing'))
    c.ok(plumb['actual'] < 11000, f'plumbing line still double counts AMP-311: {plumb["actual"]}')
    c.ok(abs(r['legacy'] - 92869.94) <= EPS, f'lineFin actual across the lines {r["legacy"]} (want $92,869.94: AMP-311 counted once)')
    c.ok(not errs, f'lines errors {errs[:3]}')
    await b.close()

async def part_overhead(p, c):
    b, pg, errs = await page(p)
    B = await books(pg)
    O = await pg.evaluate("() => { const O = rpOverhead(['2026-01-01', today()]); return {months: O.months.map(m => ({ym: m.ym, ohT: m.ohT, income: m.income, gross: m.gross, lines: m.lines})), avg: O.avgOH, gm: O.gm, be: O.breakEven, full: O.full.map(m => m.ym)} }")
    c.ok(O['full'] == ['2026-06', '2026-07', '2026-08'], f'full months {O["full"]}')
    for m in O['months']:
        M = B.pl_month(m['ym'], [x['id'] for x in B.D['projects']])
        c.ok(abs(m['ohT'] - r2(M['pool'] + M['uncoded'])) <= EPS, f"{m['ym']} overhead {m['ohT']} vs recomputed {r2(M['pool'] + M['uncoded'])}")
        c.ok(abs(sum(m['lines'].values()) - m['ohT']) <= 0.01, f"{m['ym']} categories add to {sum(m['lines'].values())}, not {m['ohT']}")
    full = [m for m in O['months'] if m['ym'] in O['full']]
    avg = r2(sum(m['ohT'] for m in full) / 3); gm = sum(m['gross'] for m in full) / sum(m['income'] for m in full)
    c.ok(abs(O['avg'] - avg) <= EPS and abs(O['be'] - r2(avg / gm)) <= EPS, f'break-even {O["be"]} vs {r2(avg / gm)} (avg {avg}, gm {gm:.4f})')
    t = await tab(pg, 'financials', 'ty')
    whole = f"${round(O['be']):,}"
    c.ok('OVERHEAD ANALYSIS' in t.upper() and whole in t, f'break-even {whole} not on Statements')
    c.ok('ACCOUNTING CHOICES TO SETTLE WITH THE ACCOUNTANT' in t.upper() and t.count('default') >= 3, 'the accounting choices are not shown with their defaults')
    c.ok(not errs, f'overhead errors {errs[:3]}')
    await b.close()

async def part_cash13(p, c):
    b, pg, errs = await page(p)
    F = await pg.evaluate("() => { const F = cashForecast(), B = companyBalance(); return {cash: F.cash, low: F.low, bank: B.cash, und: B.undeposited, weeks: F.weeks.map(w => ({net: w.net, balance: w.balance, items: w.items}))} }")
    items = [i[0] for w in F['weeks'] for i in w['items']]
    c.ok(abs(F['cash'] - r2(F['bank'] + F['und'])) <= EPS, f'cash today {F["cash"]} = bank + undeposited')
    c.ok(abs(r2(F['cash'] + sum(w['net'] for w in F['weeks'])) - F['weeks'][-1]['balance']) <= 0.01, 'cash today + the weeks\' net ≠ the last balance')
    c.ok(abs(F['low'] - min(w['balance'] for w in F['weeks'])) <= EPS, '13-week low')
    c.ok(not any(x.startswith('Bill Window Works') or x.startswith('Bill Alberto') for x in items), f'the forecast pays a bill the bank already paid: {[x for x in items if x.startswith("Bill")]}')
    c.ok(any('Payroll 8/17' in x and 'owed' in x for x in items), 'the owed payroll run is not in the forecast')
    c.ok(sum(1 for x in items if x.startswith('Card balance')) == 2, 'card balances are not in the forecast')
    t = await tab(pg, 'cash13')
    low = await pg.evaluate("() => +document.querySelector('[data-rplow]').dataset.rplow")
    c.ok(abs(low - F['low']) <= EPS and '13-WEEK CASH FLOW' in t.upper() and 'BACKLOG' in t.upper() and 'RETAINAGE' in t.upper(), 'Cash Forecast tab')
    c.ok(not errs, f'cash13 errors {errs[:3]}')
    await b.close()

async def part_excel(p, c):
    b, pg, errs = await page(p)
    await tab(pg, 'company', 'ty')
    r = await pg.evaluate("""() => { const sh = rpScreenSheets(); const u8 = rpXlsxBytes(sh); let s = ''; u8.forEach(x => s += String.fromCharCode(x));
        document.getElementById('rep-xls').click(); return {names: sh.map(x => x.name), b64: btoa(s), last: RP.lastExport, toast: document.getElementById('toast').innerText} }""")
    import openpyxl
    wb = openpyxl.load_workbook(io.BytesIO(base64.b64decode(r['b64'])))
    c.ok(len(wb.sheetnames) == len(r['names']) >= 4, f'sheets {wb.sheetnames}')
    ws = next((wb[n] for n in wb.sheetnames if n.upper().startswith('P&L BY PROJECT')), None)
    c.ok(ws is not None, f'no P&L by Project sheet: {wb.sheetnames}')
    if ws:
        vals = [cell.value for row in ws.iter_rows() for cell in row]
        c.ok(20191.11 in vals and -5329.88 in vals, f'the P&L by Project sheet lacks the numbers as numbers: {vals[:30]}')
    c.ok(r['last'] and r['last']['file'].endswith('.xlsx') and 'Exported' in r['toast'], f'Excel button: {r["last"]} {r["toast"]!r}')
    c.ok(not errs, f'excel errors {errs[:3]}')
    await b.close()

async def ask(pg, q):
    if await pg.evaluate('view') != 'pa' or not await pg.evaluate("() => !!document.getElementById('pa-in')"):
        await pg.evaluate("() => { const b = document.getElementById('tb-ask'); b ? b.click() : (view = 'pa', render()) }")
        await pg.wait_for_timeout(250)
    n = await pg.evaluate("() => document.querySelectorAll('#v-pa .pa-m.me').length")
    await pg.fill('#pa-in', q)
    await pg.press('#pa-in', 'Enter')
    try:
        await pg.wait_for_function("""n => { const Q = document.querySelectorAll('#v-pa .pa-m.me'); if (Q.length <= n) return false;
            const a = Q[Q.length - 1].nextElementSibling; const b = a && a.matches('.pa-m.a') && a.querySelector('.pa-b');
            return b && b.innerText.trim() && !/Reading the records/i.test(b.innerText) }""", arg=n, timeout=20000)
    except Exception:
        return ''
    return await pg.evaluate("() => { const Q = [...document.querySelectorAll('#v-pa .pa-m.me')].pop(); const a = Q && Q.nextElementSibling; const b = a && a.querySelector('.pa-b'); return b ? b.innerText.trim() : '' }")

async def part_ask(p, c):
    b, pg, errs = await page(p)
    W = await pg.evaluate("""() => { const S = rpPLByClient(rpRange()), O = rpOverhead(rpRange()), F = cashForecast(), B = companyBalance();
        return {net: S.net, rows: S.rows.map(r => [rpProjName(r.id), r.net]), clients: S.clients.map(x => [x.name, x.net]), be: O.breakEven, low: F.low, ret: B.ret, rp: fin2RetainagePayable().total || 0} }""")
    a = await ask(pg, 'P&L by project year to date')
    c.ok(money(W['net']) in a and all(money(n) in a for _, n in W['rows']), f'Ask P&L by project: {a[:300]!r}')
    a = await ask(pg, 'net profit by client')
    c.ok(all(money(n) in a for _, n in W['clients']), f'Ask P&L by client: {a[:300]!r}')
    a = await ask(pg, 'what is our break-even revenue')
    c.ok(money(W['be']) in a, f'Ask break-even ({money(W["be"])}): {a[:200]!r}')
    a = await ask(pg, 'what is the 13-week cash low')
    c.ok(money(W['low']) in a, f'Ask 13-week low ({money(W["low"])}): {a[:200]!r}')
    a = await ask(pg, 'retainage receivable')
    c.ok(money(W['ret']) in a and 'receivable' in a, f'Ask retainage receivable ({money(W["ret"])}): {a[:200]!r}')
    a = await ask(pg, 'how much retainage payable do we hold from subs')
    c.ok(money(W['rp']) in a and 'payable' in a, f'Ask retainage payable: {a[:200]!r}')
    a = await ask(pg, 'do the reports tie')
    c.ok(re.search(r'All \d+ report ties tie to the cent', a), f'Ask do the reports tie: {a[:200]!r}')
    # the tables: money, for a role that opens Reports
    for uid, want in (('u_owner', True), ('u_controller', True), ('u_billing', True), ('u_angel', False), ('u_alberto', False), ('u_klein', False)):
        await sign_in(pg, uid)
        r = await pg.evaluate("() => ['pnlbyproject', 'pnlbyclient', 'reportties', 'cashforecast', 'retainage', 'overheadmonths', 'linecost'].map(e => { const q = erpQuery({entity: e}); return [e, (q.rows || []).length, q.error || ''] })")
        if want: c.ok(all(n > 0 for _, n, _e in r), f'{uid} cannot read the report tables: {r}')
        else: c.ok(all(n == 0 for _, n, _e in r), f'{uid} reads money report tables: {r}')
    await sign_in(pg, 'u_owner')
    q = await pg.evaluate("() => erpQuery({entity: 'pnlbyproject'}).rows.find(r => r.projectId === 'p121')")
    c.ok(q and abs(q['net'] - dict(W['rows'])['121 State Street']) <= EPS, f'pnlbyproject 121 State net {q}')
    c.ok(not errs, f'ask errors {errs[:3]}')
    await b.close()

async def part_phone(p, c):
    for uid in ('u_owner', 'u_controller'):
        b, pg, errs = await page(p, uid, phone=True)
        for tb in ('company', 'financials', 'pack', 'profit', 'cash13', 'wip', 'ar'):
            await tab(pg, tb, 'ty', 'p121')
            r = await pg.evaluate("""() => { const sw = document.documentElement.scrollWidth, L = [...document.querySelectorAll('.rp-tie-list')];
                const cols = L.map(l => getComputedStyle(l).gridTemplateColumns.split(' ').length);
                const A = [...document.querySelectorAll('.rp-part a.rp-amt')].filter(a => a.offsetParent), taps = A.map(a => a.getBoundingClientRect().height);
                let gap = 99;
                for (let i = 1; i < A.length; i++) { const a = A[i - 1].getBoundingClientRect(), b = A[i].getBoundingClientRect();
                    if (b.top >= a.bottom - 1 && Math.abs(b.left - a.left) < 200) gap = Math.min(gap, b.top - a.bottom); }
                return {sw, cols, minTap: taps.length ? Math.min(...taps) : 99, gap} }""")
            c.ok(r['sw'] <= 390, f'{uid} phone {tb}: page scrolls sideways ({r["sw"]}px)')
            c.ok(all(x == 1 for x in r['cols']), f'{uid} phone {tb}: ties not one column {r["cols"]}')
            c.ok(r['minTap'] >= 44, f'{uid} phone {tb}: a tie amount is a {r["minTap"]}px target')
            c.ok(r['gap'] >= 8, f'{uid} phone {tb}: tie amounts {r["gap"]}px apart (the screen standard wants 8)')
        c.ok(not errs, f'{uid} phone errors {errs[:3]}')
        await b.close()

async def part_closeout(p, c):
    b, pg, errs = await page(p)
    # not in close-out: the review shows, the price-book button does not
    t = await tab(pg, 'profit', 'ty', 'p121')
    c.ok('ESTIMATE VERSUS ACTUAL' in t.upper() and 'open at close-out' in t, f'121 State: no estimate versus actual before close-out: {t[t.upper().find("ESTIMATE VERSUS"):][:200]!r}')
    c.ok(not await pg.evaluate("() => document.querySelectorAll('[data-rpcostbook]').length"), 'a price-book button before close-out')
    # 45 Cranberry reviewed to completion (the PM's cost-to-complete review, nothing left to spend): in close-out
    await pg.evaluate("() => { const p = DB.projects.find(x => x.id === 'p45'); p.ctcReviewed = today(); p.ctcBy = 'Alan'; p.ctc = 0; save(); }")
    t = await tab(pg, 'profit', 'ty', 'p45')
    rows = await pg.evaluate("() => rpCloseout(DB.projects.find(x => x.id === 'p45')).rows")
    arch = next((r for r in rows if r['code'] == '15'), None)
    c.ok(arch and abs(arch['final'] - 7032) < EPS and abs(arch['variance'] - r2(7032 - arch['budget'])) < EPS and arch['proposable'], f'45 Cranberry code 15 at close-out: {arch}')
    n0 = await pg.evaluate("() => (DB.changes || []).length")
    await pg.evaluate("() => document.querySelector('[data-rpcostbook=\"p45:15\"]').click()")
    await pg.wait_for_timeout(200)
    r = await pg.evaluate("() => { const ch = DB.changes.slice(-1)[0]; return {n: DB.changes.length, ch, toast: document.getElementById('toast').innerText, row: document.querySelector('[data-rpcloseout]').innerText} }")
    ch = r['ch']
    c.ok(r['n'] == n0 + 1 and ch['status'] == 'draft' and ch['delta']['path'] == 'company.costBook.15' and abs(ch['delta']['to'] - r2(7032 / arch['budget'])) < EPS, f'the proposal: {ch}')
    c.ok('Estimate vs. actual review at close-out'.lower() in ch['why'].lower() and 'Every dollar matters' in ch['why'], f'the proposal does not name its Standard: {ch["why"]!r}')
    c.ok('Proposed to the price book' in r['toast'] and 'Open' in r['toast'], f'toast {r["toast"]!r}')
    c.ok('Proposed, awaiting the owner' in r['row'], 'the row does not show the proposal')
    # the toast's Open lands on Company › Changes; the owner approves; the cost factor changes only then
    c.ok(await pg.evaluate("() => (DB.company.costBook || {})['15'] === undefined"), 'the cost factor changed before approval')
    await pg.evaluate("() => { const a = [...document.querySelectorAll('#toast a')].find(x => /Open/.test(x.innerText)); a && a.click() }")
    await pg.wait_for_timeout(250)
    v = await pg.evaluate("() => [view, typeof mtTab !== 'undefined' ? mtTab : '']")
    c.ok(v == ['meetings', 'changes'], f'Open went to {v}')
    # proposing twice is refused
    await pg.evaluate("() => rpProposeCostBook('p45', '15')")
    c.ok(await pg.evaluate("() => DB.changes.filter(x => x.key === 'costbook:p45:15').length") == 1, 'proposed twice')
    await pg.evaluate("() => { view = 'meetings'; mtTab = 'changes'; mtOpen = ''; render(); }")
    await pg.wait_for_timeout(150)
    ok = await pg.evaluate("() => { const id = DB.changes.find(x => x.key === 'costbook:p45:15').id; const btn = document.querySelector(`[data-chok=\"${id}\"]`); if (!btn) return 'no approve button'; btn.click(); return (DB.company.costBook || {})['15'] }")
    c.ok(ok == ch['delta']['to'], f'approving the change did not set the cost factor: {ok}')
    t = await tab(pg, 'profit', 'ty', 'p45')
    c.ok('In the price book' in t, 'the row does not say the code is in the price book')
    c.ok(not errs, f'closeout errors {errs[:3]}')
    await b.close()

async def part_manual(p, c):
    """A manual journal, drafted by the controller and posted by the owner with a reason, is flagged on the ties
    with who, when and why, and its drill marks every line as posted by hand. It moves no tie: the management view
    carries nothing posted by hand, the bridge to the ledger names it."""
    b, pg, errs = await page(p, 'u_controller')
    await pg.evaluate("() => { glTab = 'entry'; glEnt = 'all'; view = 'gl'; render(); }")
    await pg.wait_for_timeout(150)
    await pg.fill('#gle-date', '2026-08-31')
    await pg.fill('#gle-memo', 'August depreciation on the van')
    await pg.evaluate("""() => { const rows = document.querySelectorAll('#gle-lines tbody tr');
        rows[0].querySelector('[data-l=acct]').value = '8000'; rows[0].querySelector('[data-l=dr]').value = '1200';
        rows[1].querySelector('[data-l=acct]').value = '1590'; rows[1].querySelector('[data-l=cr]').value = '1200'; }""")
    await pg.click('#gle-save')
    await pg.wait_for_timeout(200)
    await sign_in(pg, 'u_owner')
    st = await pg.evaluate("() => { const j = DB.gl.journals.find(x => x.kind === 'manual' && /van/.test(x.memo)); try { mbCloseModal(); } catch (e) {} glApprove(j.id, 'Straight-line on the van, August'); return j.status }")
    c.ok(st == 'posted', f'the manual journal posted: {st}')
    await tab(pg, 'financials', 'ty')
    T = await screen_ties(pg, '[data-rpties=all]')
    m = [x for x in T if x['manual']]
    c.ok(len(m) == 1 and 'Flagged' in m[0]['text'] and 'Straight-line on the van' in m[0]['text'] and '$1,200.00' in m[0]['text'], f'the manual journal is not flagged on the ties: {m}')
    c.ok(all(x['ok'] for x in T if not x['manual']), f'a manual journal moved a tie: {[x["id"] for x in T if not x["manual"] and not x["ok"]]}')
    await pg.evaluate("() => document.querySelector('[data-rptie=manual] [data-rpdrill]').click()")
    await pg.wait_for_timeout(200)
    t = await pg.evaluate("() => document.getElementById('mbm').innerText")
    c.ok(t.count('By Hand') >= 2 and 'posted by hand' in t, f'the manual journal drill does not mark its lines: {t[:300]!r}')
    await pg.evaluate("() => mbCloseModal()")
    t = await tab(pg, 'pack', 'lm')
    c.ok('Flagged' in t and 'Straight-line on the van' in t, 'the Monthly Pack does not flag the manual journal on its ties')
    c.ok(not errs, f'manual errors {errs[:3]}')
    await b.close()

PARTS = {'books': part_books, 'screen': part_screen, 'cash': part_cash, 'sealed': part_sealed, 'entity': part_entity, 'lines': part_lines,
         'overhead': part_overhead, 'closeout': part_closeout, 'manual': part_manual, 'cash13': part_cash13, 'excel': part_excel, 'ask': part_ask, 'phone': part_phone}

async def main():
    want = [a for a in sys.argv[1:] if a in PARTS] or list(PARTS)
    c = Checks()
    async with async_playwright() as p:
        for k in want:
            n0, f0 = c.n, len(c.fail)
            await PARTS[k](p, c)
            print(f'-- {k}: {c.n - n0} checks, {len(c.fail) - f0} failed')
    sys.exit(c.done())

asyncio.run(main())
